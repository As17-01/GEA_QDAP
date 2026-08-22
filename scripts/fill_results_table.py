#!/usr/bin/env python3
"""
Fill the Table_of_Results ALGEA/GQAP Excel workbook from scripts/results/*.json.

Experiment 1: standard-family results for the original/no-repair setup.
Experiment 2: improved-family results for the new chromosome/repair setup.

Usage:
    python scripts/fill_results_table.py
    python scripts/fill_results_table.py --xlsx /path/to/Table_of_Results_ALGEA_GQAP.xlsx
"""

from __future__ import annotations

import argparse
import math
import sys
from pathlib import Path

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from build_results_table import _get_stats, load_results  # noqa: E402

DEFAULT_XLSX = Path(
    "/mnt/c/Users/Aleksej Seliverstov/Downloads/Table_of_Results_ALGEA_GQAP.xlsx"
)

# (results/*.json stem, first column for Mean ± Std)
IMPROVED_ALGO_COLUMNS: list[tuple[str, int]] = [
    ("ga", 2),
    ("gea_scenario_1", 8),
    ("gea_scenario_2", 14),
    ("gea_scenario_3", 20),
    ("gea", 26),
    ("improvedsa", 32),
    ("improvedparticleswarm", 38),
    ("improvedhybridgasa", 44),
    ("improvedhybridgapso", 50),
    ("adaptive", 56),
    ("adaptive_gea_scenario_1", 62),
    ("adaptive_gea_scenario_2", 68),
    ("adaptive_gea_scenario_3", 74),
    ("adaptive_gea", 80),
]

STANDARD_ALGO_COLUMNS: list[tuple[str, int]] = [
    ("standard", 2),
    ("standard_gea_scenario_1", 8),
    ("standard_gea_scenario_2", 14),
    ("standard_gea_scenario_3", 20),
    ("standard_gea", 26),
    ("standard_sa", 32),
    ("standard_pso", 38),
    ("standard_hybrid_ga_sa", 44),
    ("standard_hybrid_ga_pso", 50),
    ("standard_adaptive", 56),
    ("standard_adaptive_gea_scenario_1", 62),
    ("standard_adaptive_gea_scenario_2", 68),
    ("standard_adaptive_gea_scenario_3", 74),
    ("standard_adaptive_gea", 80),
]

DATA_START_ROW = 5
DATA_END_ROW = 33


def _fmt(v: float | None, decimals: int = 0) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if decimals == 0:
        return f"{v:,.0f}"
    return f"{v:,.{decimals}f}"


def _record_cells(record: dict) -> tuple[str, str, str, str, str]:
    """Return (mean±std, best, worst, hit, nfe) strings for one dataset record."""
    r = _get_stats(record, "results") or {}
    ht = _get_stats(record, "hitting_time") or {}
    nfe = _get_stats(record, "nfe") or {}

    mean_ = r.get("mean")
    std_ = r.get("std")
    min_ = r.get("min")
    max_ = r.get("max")
    hit_ = ht.get("mean")
    nfe_ = nfe.get("mean")

    if mean_ is not None and std_ is not None:
        mean_std = f"{_fmt(mean_)} ± {_fmt(std_)}"
    else:
        mean_std = ""

    return (
        mean_std,
        _fmt(min_),
        _fmt(max_),
        _fmt(hit_, 1),
        _fmt(nfe_),
    )


def _dataset_rows(ws) -> dict[str, int]:
    """Map dataset name → Excel row number (column A, rows 5–33)."""
    mapping: dict[str, int] = {}
    for row in range(DATA_START_ROW, DATA_END_ROW + 1):
        name = ws.cell(row, 1).value
        if name:
            mapping[str(name)] = row
    return mapping


def fill_sheet(
    ws,
    data: dict[str, dict[str, dict]],
    algo_columns: list[tuple[str, int]],
) -> int:
    """Write result values into ws; return number of cells written."""
    dataset_rows = _dataset_rows(ws)
    written = 0

    for stem, start_col in algo_columns:
        if stem not in data:
            print(f"  Warning: no results for {stem!r}, skipping column {start_col}")
            continue

        for dataset, row in dataset_rows.items():
            if dataset not in data[stem]:
                continue
            mean_std, best, worst, hit, nfe = _record_cells(data[stem][dataset])
            values = (mean_std, best, worst, hit, nfe)
            for offset, value in enumerate(values):
                if value:
                    ws.cell(row, start_col + offset, value)
                    written += 1

    return written


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--results-dir",
        type=Path,
        default=SCRIPT_DIR / "results",
        help="Directory with per-algorithm JSON files (default: scripts/results/)",
    )
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=DEFAULT_XLSX,
        help="Path to Table_of_Results ALGEA/GQAP workbook",
    )
    args = parser.parse_args()

    if not args.xlsx.is_file():
        print(f"Workbook not found: {args.xlsx}")
        sys.exit(1)

    print(f"Loading results from {args.results_dir} ...")
    data = load_results(args.results_dir)
    if not data:
        print("No result JSON files found.")
        sys.exit(1)

    print(f"Opening {args.xlsx} ...")
    wb = load_workbook(args.xlsx)

    if "Experiment 2" not in wb.sheetnames:
        print("Sheet 'Experiment 2' not found.")
        sys.exit(1)

    ws2 = wb["Experiment 2"]
    n2 = fill_sheet(ws2, data, IMPROVED_ALGO_COLUMNS)
    print(f"Experiment 2 (improved): wrote {n2} cells")

    if "Experiment 1" in wb.sheetnames:
        ws1 = wb["Experiment 1"]
        n1 = fill_sheet(ws1, data, STANDARD_ALGO_COLUMNS)
        print(f"Experiment 1 (standard): wrote {n1} cells")
    else:
        print("Sheet 'Experiment 1' not found, skipping.")

    wb.save(args.xlsx)
    print(f"Saved {args.xlsx}")


if __name__ == "__main__":
    main()
